"""Smart Import files, matching, transactions and retries on a generated PostgreSQL DB only."""
import base64
import importlib.util
import io
import json
import sys
import unittest
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
spec = importlib.util.spec_from_file_location('smart_fixture', Path(__file__).resolve().parents[1] / 'frontend/tests/phase3-api.py')
fixture = importlib.util.module_from_spec(spec); spec.loader.exec_module(fixture)
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from backend.services.smart_import_files import parse_file, template, auto_mapping


def xlsx(rows):
    workbook = Workbook(); sheet = workbook.active
    for row in rows: sheet.append(row)
    output = io.BytesIO(); workbook.save(output)
    return base64.b64encode(output.getvalue()).decode()


class SpreadsheetTests(unittest.TestCase):
    def test_templates(self):
        for mode in ['customers','assets','combined']:
            for format in ['csv','xlsx']:
                result = template(mode, format)
                content = base64.b64decode(result['content'])
                if format == 'csv': headings = content.decode('utf-8-sig').splitlines()[0]
                else:
                    workbook = load_workbook(io.BytesIO(content)); headings = str([c.value for c in workbook.active[1]])
                    self.assertEqual(workbook.active['B2'].number_format,'@'); workbook.close()
                self.assertIn('Customer Name',headings); self.assertNotIn('customer_reference',headings); self.assertNotIn('Warranty End',headings)
    def test_csv_xlsx_dates_and_phone_zeroes(self):
        from datetime import date
        result=parse_file(xlsx([['Customer Name','Phone','Installation Date'],['Ahmed','01012345678',date(2024,2,29)]]),'xlsx')
        self.assertEqual(result['rows'][0],['Ahmed','01012345678','2024-02-29'])
        source=base64.b64encode('اسم العميل,الموبايل,نوع الجهاز,المحافظة,المنطقة,Unknown\nAhmed,01012345678,Filter,Cairo,Nasr City,x\n'.encode()).decode()
        result=parse_file(source,'csv'); self.assertEqual(result['mapping'],['customer_name','phone','asset_type','state','area',''])
    def test_duplicate_aliases_require_manual_mapping(self):
        self.assertEqual(auto_mapping(['Client Name','Phone','Mobile']),['customer_name','phone',''])
    def test_malformed_limits_and_content(self):
        values=[('xlsx',base64.b64encode(b'not a workbook').decode()),('csv',base64.b64encode(b'Name,Phone\n"unclosed').decode()),('csv',base64.b64encode(b'Name,Phone\na,'+b'x'*1001).decode()),('csv',base64.b64encode(b'Name,Phone\n'+b'a,b\n'*1001).decode()),('csv',base64.b64encode(b'x'*1048577).decode()),('csv',base64.b64encode(b'Name,Phone\na,=1+1').decode()),('csv',base64.b64encode(b'Name,Name\na,b').decode()),('csv',base64.b64encode(b'Name,Phone\na,\xff').decode())]
        for format, content in values:
            with self.subTest(format=format,content=content[:20]),self.assertRaises(ValueError): parse_file(content,format)
    def test_formula_external_links_macros_and_xml_entities(self):
        with self.assertRaises(ValueError): parse_file(xlsx([['Name','Phone'],['=1+1','01012345678']]),'xlsx')
        clean=base64.b64decode(xlsx([['Name','Phone'],['Ahmed','01012345678']]))
        for filename, content in [('xl/externalLinks/externalLink1.xml',b'<x/>'),('xl/vbaProject.bin',b'not executable'),('macro.xml',b'<Override ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml"/>'),('evil.xml',b'<!DOCTYPE x [<!ENTITY x "boom">]><x>&x;</x>'),('xl/worksheets/_rels/sheet1.xml.rels',b'<Relationships><Relationship TargetMode="External" Target="https://example.test"/></Relationships>')]:
            with io.BytesIO(clean) as output:
                with zipfile.ZipFile(output,'a') as archive: archive.writestr(filename,content)
                with self.subTest(filename=filename),self.assertRaises(ValueError): parse_file(base64.b64encode(output.getvalue()).decode(),'xlsx')
    def test_expanded_zip_and_sparse_row_limits(self):
        for content in [b'x'*(17*1048576),b'<worksheet><row r="10000000"/></worksheet>']:
            output=io.BytesIO()
            with zipfile.ZipFile(output,'w',compression=zipfile.ZIP_DEFLATED) as archive: archive.writestr('xl/worksheets/sheet1.xml',content)
            with self.assertRaises(ValueError): parse_file(base64.b64encode(output.getvalue()).decode(),'xlsx')


class SmartImportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fixture=fixture.synthetic_application(); cls.client=TestClient(cls.fixture.__enter__()); cls.headers={}
        for role in ['admin','foreign','technician']:
            response=cls.client.post('/api/v1/auth/login',data={'username':role+'@example.test','password':'synthetic-password'})
            cls.headers[role]={'Authorization':'Bearer '+response.json()['access_token']}
    @classmethod
    def tearDownClass(cls): cls.client.close(); cls.fixture.__exit__(None,None,None)
    def data(self,mode='combined',phone='01012345678',name=None,serial=None):
        return {'mode':mode,'columns':['Customer Name','Phone','Asset Type','Serial Number','Country','State','Installation Date','Maintenance Cycle','Warranty Years'], 'mapping':['customer_name','phone','asset_type','serial_number','country','state','installation_date','maintenance_cycle','warranty_years'], 'rows':[[name or 'Smart '+uuid4().hex,phone,'Filter',serial or uuid4().hex,' Egypt ',' cairo ','2024-02-29','6','1']], 'corrections':{}}
    def preview(self,data,role='admin'):
        return self.client.post('/api/v1/imports/preview',headers=self.headers[role],json=data)
    def commit(self,data,preview=None,key=None,role='admin'):
        preview=preview or self.preview(data,role).json()
        return self.client.post('/api/v1/imports/commit',headers=self.headers[role],json={'data':data,'token':preview.get('token',''),'idempotency_key':key or str(uuid4())})
    def customer(self,name,phone,role='admin'):
        return self.client.post('/api/v1/customers',headers=self.headers[role],json={'name':name,'phones':[{'country':'EG','number':phone,'label':'Primary'}]}).json()
    def test_01_combined_grouping_and_warranty(self):
        data=self.data();data['rows'].append([*data['rows'][0]]);data['rows'][1][3]=uuid4().hex
        preview=self.preview(data).json();self.assertTrue(preview['valid'],preview);self.assertEqual(preview['summary']['new_customers'],1);self.assertEqual(preview['summary']['new_assets'],2)
        self.assertEqual(preview['preview'][0]['warranty_end'],'2025-02-28')
        result=self.commit(data,preview).json();self.assertEqual(len(result['customer_numbers']),1);self.assertEqual(len(result['asset_numbers']),2)
        customer=next(c for c in self.client.get('/api/v1/customers',headers=self.headers['admin']).json() if c['display_id']==result['customer_numbers'][0])
        assets=[a for a in self.client.get('/api/v1/assets',headers=self.headers['admin']).json() if a['customer_id']==customer['id']]
        self.assertEqual(len(assets),2);self.assertEqual(assets[0]['country'],'EG');self.assertEqual(assets[0]['state'],'EG-C');self.assertEqual(customer['phones'][0]['normalized'],'+201012345678')
    def test_02_existing_phone_customers_only_and_assets_only(self):
        for mode in ['customers','assets','combined']:
            data=self.data(mode,phone='+201000000001');preview=self.preview(data).json();self.assertTrue(preview['valid'],preview)
            self.assertEqual(preview['summary']['new_customers'],0);self.assertEqual(preview['summary']['existing_customers'],1)
            result=self.commit(data,preview).json();self.assertEqual(result['customer_numbers'],[]);self.assertEqual(len(result['asset_numbers']),0 if mode=='customers' else 1)
        data=self.data('customers',phone='01000000003');result=self.commit(data).json();self.assertEqual(len(result['customer_numbers']),1);self.assertEqual(result['asset_numbers'],[])
    def test_03_ambiguous_phone_and_explicit_customer_correction(self):
        a=self.customer('Ambiguous A','01000000004');self.customer('Ambiguous B','01000000004')
        data=self.data(phone='01000000004');preview=self.preview(data).json();self.assertFalse(preview['valid']);self.assertIn('Multiple customers',str(preview['errors']))
        data['corrections']={'2':{'customer_selection':str(a['display_id'])}}
        preview=self.preview(data).json();self.assertTrue(preview['valid'],preview);self.assertEqual(preview['summary']['new_customers'],0)
        self.assertTrue(self.commit(data,preview).json()['valid'])
    def test_04_foreign_phone_cannot_match_and_cross_tenant_token(self):
        foreign=self.customer('Other tenant','01000000005','foreign');data=self.data(phone='01000000005')
        preview=self.preview(data).json();self.assertEqual(preview['summary']['existing_customers'],0);self.assertEqual(preview['summary']['new_customers'],1)
        self.assertNotIn(foreign['id'],json.dumps(preview));self.assertEqual(self.commit(data,preview,role='foreign').status_code,400)
        data['mode']='assets';self.assertFalse(self.preview(data).json()['valid'])
    def test_05_geography_errors_and_corrections(self):
        data=self.data();data['rows'][0][5]='Cairoo';preview=self.preview(data).json();self.assertFalse(preview['valid']);self.assertEqual(preview['errors'][0]['field'],'state');self.assertIn('Cairoo',preview['errors'][0]['message'])
        data['corrections']={'2':{'state':'EG-C'}};self.assertTrue(self.preview(data).json()['valid'])
        data['rows'][0][4]='Atlantis';self.assertFalse(self.preview(data).json()['valid'])
        from backend.services.smart_import import resolve_state,resolve_country
        self.assertEqual(resolve_state('US',' California '),'US-CA');self.assertEqual(resolve_country(' egypt '),'EG')
        with patch('backend.services.smart_import.geography',return_value={'countries':[], 'states':{'US':[{'code':'US-CA','name':'Same'},{'code':'US-TX','name':'Same'}]}}):
            with self.assertRaises(ValueError): resolve_state('US','Same')
    def test_06_conflicting_new_group_requires_attention(self):
        data=self.data(phone='01000000006');data['rows'].append([*data['rows'][0]]);data['rows'][1][0]='Different name';data['rows'][1][3]=uuid4().hex
        preview=self.preview(data).json();self.assertFalse(preview['valid']);self.assertEqual(preview['summary']['attention_rows'],2)
        data['corrections']={'3':{'customer_name':data['rows'][0][0]}};self.assertTrue(self.preview(data).json()['valid'])
    def test_07_serial_collision_deleted_reservation_and_case(self):
        data=self.data(phone='+201000000001',serial='SYN-001');self.assertIn('duplicate_tenant',[e['code'] for e in self.preview(data).json()['errors']])
        data['rows'][0][3]='syn-001';self.assertTrue(self.preview(data).json()['valid'])
        data=self.data(phone='+201000000001');data['rows'][0][3]='  '+data['rows'][0][3]+'  ';self.commit(data)
        asset=next(a for a in self.client.get('/api/v1/assets',headers=self.headers['admin']).json() if a['serial_number']==data['rows'][0][3])
        self.client.delete('/api/v1/assets/'+asset['id'],headers=self.headers['admin']);self.assertFalse(self.preview(data).json()['valid'])
        data['rows'].append([*data['rows'][0]]);self.assertIn('duplicate_file',[e['code'] for e in self.preview(data).json()['errors']])
    def test_08_service_validation_and_ambiguous_dates(self):
        for index,value,field in [(7,'1.5','maintenance_cycle'),(7,'0','maintenance_cycle'),(8,'1.5','warranty_years'),(8,'101','warranty_years'),(6,'01/02/2026','installation_date')]:
            data=self.data();data['rows'][0][index]=value;result=self.preview(data).json();self.assertFalse(result['valid']);self.assertIn(field,[e['field'] for e in result['errors']])
    def test_09_idempotency_expiry_and_changed_input(self):
        import jwt
        from backend.core.config import get_settings
        data=self.data(phone='+201000000001');preview=self.preview(data).json();key=str(uuid4())
        first=self.commit(data,preview,key).json();again=self.commit(data,preview,key).json();self.assertTrue(again['replayed']);self.assertEqual(first['asset_numbers'],again['asset_numbers'])
        claims=jwt.decode(preview['token'],get_settings().secret_key,algorithms=['HS256']);claims['exp']=1
        expired={'token':jwt.encode(claims,get_settings().secret_key,algorithm='HS256')};self.assertTrue(self.commit(data,expired,key).json()['replayed'])
        self.assertEqual(self.commit(data,expired).status_code,400)
        data['rows'][0][3]=uuid4().hex;self.assertEqual(self.commit(data,preview,key).status_code,400)
    def test_10_atomic_rollback_and_retry(self):
        from backend.services.asset import AssetService
        data=self.data(phone='01000000007');data['rows'].append([*data['rows'][0]]);data['rows'][1][3]=uuid4().hex;preview=self.preview(data).json();key=str(uuid4())
        before=self.client.get('/api/v1/customers',headers=self.headers['admin']).json()
        original=AssetService.create_asset;calls=[]
        def fail(service,*args):
            calls.append(1)
            if len(calls)==2: raise RuntimeError('synthetic failure')
            return original(service,*args)
        with patch.object(AssetService,'create_asset',fail):
            with self.assertRaises(RuntimeError): self.commit(data,preview,key)
        self.assertEqual(before,self.client.get('/api/v1/customers',headers=self.headers['admin']).json())
        self.assertTrue(self.commit(data,preview,key).json()['valid'])
    def test_11_signed_match_plan_revalidated(self):
        data=self.data(phone='01000000008');preview=self.preview(data).json();self.customer('Created after preview','01000000008')
        self.assertEqual(self.commit(data,preview).status_code,400)
    def test_12_mapping_ignored_columns_and_limits(self):
        data=self.data();data['columns'].append('Ignore');data['mapping'].append('');data['rows'][0].append('unused');self.assertTrue(self.preview(data).json()['valid'])
        data['mapping'][-1]='phone';self.assertEqual(self.preview(data).status_code,400)
        data=self.data();data['rows']*=1001;self.assertEqual(self.preview(data).status_code,400)
        self.assertEqual(self.client.post('/api/v1/imports/upload',headers=self.headers['admin'],content='x'*(7*1048576+1)).status_code,413)
    def test_13_api_upload_templates_and_authorization(self):
        for format in ['csv','xlsx']:
            result=self.client.get('/api/v1/imports/template',params={'mode':'combined','format':format},headers=self.headers['admin']);self.assertEqual(result.status_code,200)
            content=xlsx([['Name','Phone'],['Ahmed','01012345678']]) if format=='xlsx' else base64.b64encode(b'Name,Phone\nAhmed,01012345678').decode()
            uploaded=self.client.post('/api/v1/imports/upload',headers=self.headers['admin'],json={'mode':'customers','format':format,'content':content});self.assertEqual(uploaded.status_code,200);self.assertEqual(uploaded.json()['mapping'],['customer_name','phone'])
        self.assertEqual(self.client.post('/api/v1/imports/preview',json=self.data()).status_code,401)
        self.assertEqual(self.preview(self.data(),'technician').status_code,403)
        self.assertEqual(self.client.post('/api/v1/imports/upload',headers=self.headers['technician'],json={'mode':'customers'}).status_code,403)
    def test_14_concurrent_same_phone_group_and_same_key(self):
        from backend.core.database import SessionLocal
        from backend.models.user import User
        from sqlalchemy import select
        from backend.services.smart_import import commit
        with SessionLocal() as db:
            user=db.scalar(select(User).where(User.email=='admin@example.test'));company=user.company_id;actor=user.id
        data=self.data(phone='01000000009');preview=self.preview(data).json();key=str(uuid4())
        def write(_):
            with SessionLocal() as db:
                result=commit(db,company,actor,data,preview['token'],key);db.commit();return result
        with ThreadPoolExecutor(max_workers=2) as pool: results=list(pool.map(write,range(2)))
        self.assertEqual(results[0]['customer_numbers'],results[1]['customer_numbers']);self.assertEqual(sum(r['replayed'] for r in results),1)
    def test_15_concurrent_different_keys_serial_race(self):
        from backend.core.database import SessionLocal
        from backend.models.user import User
        from sqlalchemy import select
        from backend.services.smart_import import commit
        with SessionLocal() as db:
            user=db.scalar(select(User).where(User.email=='admin@example.test'));company=user.company_id;actor=user.id
        data=self.data(phone='+201000000001');preview=self.preview(data).json()
        def write(_):
            with SessionLocal() as db:
                result=commit(db,company,actor,data,preview['token'],str(uuid4()));db.commit();return result
        with ThreadPoolExecutor(max_workers=2) as pool: results=list(pool.map(write,range(2)))
        self.assertEqual(sum(r['valid'] for r in results),1)
    def test_16_explicit_selection_without_phone_and_no_name_guess(self):
        data=self.data('assets',phone='');data['corrections']={'2':{'customer_selection':'1001'}}
        self.assertTrue(self.preview(data).json()['valid']);self.assertTrue(self.commit(data).json()['valid'])
        data['corrections']={};data['rows'][0][0]='Synthetic Customer';self.assertFalse(self.preview(data).json()['valid'])
    def test_17_malformed_requests_fail_closed(self):
        for mode in [[],{},None]:
            data=self.data();data['mode']=mode;self.assertEqual(self.preview(data).status_code,400)
        self.assertEqual(self.client.post('/api/v1/imports/upload',headers=self.headers['admin'],json={'mode':'customers','format':[],'content':'a'}).status_code,400)
        data=self.data();data['corrections']={'2':{'customer_id':str(uuid4())}};self.assertEqual(self.preview(data).status_code,400)
    def test_18_existing_customer_edit_requires_fresh_review(self):
        customer=self.customer('Snapshot','01000000010');data=self.data(phone='01000000010');preview=self.preview(data).json()
        self.client.patch('/api/v1/customers/'+customer['id'],headers=self.headers['admin'],json={'name':'Changed snapshot'})
        self.assertEqual(self.commit(data,preview).status_code,400)
    def test_19_direct_service_requires_company_scope(self):
        from backend.services.smart_import import validate, commit
        with self.assertRaises(ValueError): validate(None,None,self.data())
        with self.assertRaises(ValueError): commit(None,None,uuid4(),self.data(),'token',str(uuid4()))

    def records(self, resource):
        return self.client.get('/api/v1/'+resource,headers=self.headers['admin']).json()

    def test_20_skip_collision_restore_mixed_commit_and_retry(self):
        from copy import deepcopy
        original_customers=self.records('customers');original_assets=self.records('assets')
        data=self.data(phone='+201000000001',serial='SYN-001')
        data['rows'].append(self.data(phone='01099998888',name='New Import Customer',serial='SMART-TEST-002')['rows'][0])
        initial=self.preview(data).json()
        self.assertFalse(initial['valid']);self.assertEqual(initial['preview'][0]['category'],'Needs attention')
        data['skipped_rows']=[2];review=self.preview(data).json()
        self.assertTrue(review['valid']);self.assertEqual(review['preview'][0]['category'],'Skipped')
        self.assertEqual(review['summary'],dict(new_customers=1,existing_customers=0,new_assets=1,overwritten_assets=0,attention_rows=0,skipped_rows=1))
        restored=deepcopy(data);restored['skipped_rows']=[]
        self.assertFalse(self.preview(restored).json()['valid'])
        self.assertEqual(self.commit(restored,review).status_code,400)
        key=str(uuid4());result=self.commit(data,review,key).json();again=self.commit(data,review,key).json()
        self.assertTrue(again['replayed']);self.assertEqual(result['asset_numbers'],again['asset_numbers'])
        customers=self.records('customers');assets=self.records('assets')
        for before in original_customers:self.assertEqual(before,next(c for c in customers if c['id']==before['id']))
        for before in original_assets:self.assertEqual(before,next(a for a in assets if a['id']==before['id']))
        created=[c for c in customers if c['name']=='New Import Customer'];self.assertEqual(len(created),1)
        created_assets=[a for a in assets if a['serial_number']=='SMART-TEST-002'];self.assertEqual(len(created_assets),1)
        self.assertEqual(created_assets[0]['customer_id'],created[0]['id'])

    def test_21_skip_recomputes_groups_and_conflicts(self):
        data=self.data(phone='01000000031');data['rows'].append([*data['rows'][0]])
        data['rows'][1][0]='Conflicting name';data['rows'][1][3]=uuid4().hex
        self.assertFalse(self.preview(data).json()['valid'])
        data['skipped_rows']=[2];review=self.preview(data).json()
        self.assertTrue(review['valid']);self.assertEqual(review['summary']['new_customers'],1)
        result=self.commit(data,review).json();self.assertEqual(len(result['customer_numbers']),1);self.assertEqual(len(result['asset_numbers']),1)
        self.assertFalse(any(a['serial_number']==data['rows'][0][3] for a in self.records('assets')))

    def test_22_all_skipped_has_no_token_or_receipt_even_direct_commit(self):
        from sqlalchemy import select,func
        from backend.core.database import SessionLocal
        from backend.core.config import get_settings
        from backend.models.audit_log import AuditLog
        from backend.models.user import User
        from backend.services.smart_import import preview_token
        data=self.data(phone='01000000032');data['rows'].append([*data['rows'][0]]);data['skipped_rows']=[2,3]
        before=(self.records('customers'),self.records('assets'))
        with SessionLocal() as db:
            user=db.scalar(select(User).where(User.email=='admin@example.test'))
            token=preview_token(data,user.company_id,user.id,[])
            receipts=db.scalar(select(func.count()).select_from(AuditLog))
        review=self.preview(data).json();self.assertFalse(review['valid']);self.assertNotIn('token',review)
        self.assertEqual(review['message'],'No rows selected for import.');self.assertEqual(review['summary']['new_customers'],0)
        self.assertFalse(self.commit(data,{'token':token}).json()['valid'])
        self.assertEqual(before,(self.records('customers'),self.records('assets')))
        with SessionLocal() as db:self.assertEqual(receipts,db.scalar(select(func.count()).select_from(AuditLog)))

    def test_23_skip_ready_customers_and_assets_modes(self):
        for mode,phone in [('customers','01000000033'),('assets','+201000000001')]:
            data=self.data(mode,phone=phone);data['rows'].append([*data['rows'][0]])
            data['rows'][1][1]='01000000034' if mode=='customers' else phone;data['rows'][1][3]=uuid4().hex
            data['skipped_rows']=[2];review=self.preview(data).json();self.assertTrue(review['valid'],review)
            result=self.commit(data,review).json();self.assertEqual(result['summary']['skipped_rows'],1)
            self.assertEqual(len(result['customer_numbers']),1 if mode=='customers' else 0)
            self.assertEqual(len(result['asset_numbers']),1 if mode=='assets' else 0)
            if mode=='customers':self.assertFalse(any(c['phones'] and any(p['number']==phone for p in c['phones']) for c in self.records('customers')))
            else:self.assertFalse(any(a['serial_number']==data['rows'][0][3] for a in self.records('assets')))

    def test_24_skip_business_errors_and_duplicate_file_are_excluded(self):
        data=self.data(phone='01000000035');data['rows'].append([*data['rows'][0]])
        data['rows'][0][1]='invalid';data['rows'][0][4]='Atlantis';data['rows'][0][5]='Cairoo';data['rows'][0][6]='bad date'
        self.assertFalse(self.preview(data).json()['valid'])
        data['skipped_rows']=[2];review=self.preview(data).json();self.assertTrue(review['valid'],review);self.assertEqual(review['errors'],[])
        self.assertTrue(self.commit(data,review).json()['valid'])

    def test_25_skip_state_and_source_identity_tampering(self):
        from copy import deepcopy
        data=self.data(phone='+201000000001');data['rows'].append([*data['rows'][0]]);data['rows'][1][3]=uuid4().hex
        data['skipped_rows']=[2];review=self.preview(data).json()
        for skip in [[],[3]]:
            changed=deepcopy(data);changed['skipped_rows']=skip;self.assertEqual(self.commit(changed,review).status_code,400)
        changed=deepcopy(data);changed['rows'].reverse();self.assertEqual(self.commit(changed,review).status_code,400)
        changed=deepcopy(data);changed['corrections']={'2':{'phone':'01000000036'}};self.assertEqual(self.commit(changed,review).status_code,400)

    def test_26_skip_request_bounds_and_types(self):
        for skipped in ['2',{},None,[True],['2'],[2.0],[1],[3],[2,2]]:
            data=self.data();data['skipped_rows']=skipped
            with self.subTest(skipped=skipped):self.assertEqual(self.preview(data).status_code,400)
        data=self.data();data['skipped_rows']=[2];data['rows'][0][0]='x'*1001
        self.assertEqual(self.preview(data).status_code,400)

    def test_27_skipped_foreign_phone_and_selection_expose_nothing(self):
        foreign=self.customer('Skip foreign secret','01000000037','foreign')
        data=self.data(phone='01000000037');data['rows'].append(self.data(phone='+201000000001')['rows'][0])
        data['corrections']={'2':{'customer_selection':str(foreign['display_id'])}};data['skipped_rows']=[2]
        review=self.preview(data).json();self.assertTrue(review['valid']);self.assertIsNone(review['preview'][0]['customer_number'])
        self.assertNotIn(foreign['name'],json.dumps(review));self.assertNotIn(foreign['id'],json.dumps(review))
        self.assertEqual(self.commit(data,review,role='foreign').status_code,400)
        self.assertTrue(self.commit(data,review).json()['valid'])

if __name__=='__main__': unittest.main(verbosity=2)
