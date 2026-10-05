"""Bounded spreadsheet decoding and deterministic, bilingual column matching."""
import base64
import csv
import io
import re
import zipfile
from datetime import date, datetime

from defusedxml.ElementTree import fromstring
from openpyxl import Workbook, load_workbook

MAX_BYTES = 1048576
MAX_ROWS = 1000
MAX_COLUMNS = 40
MAX_CELL = 1000
FIELDS = {
    'customer_name': ('Customer Name', 'Name', 'Client Name', 'اسم العميل'),
    'phone': ('Phone', 'Mobile', 'Phone Number', 'الموبايل', 'رقم الهاتف'),
    'phone_label': ('Phone Label', 'Label', 'وصف الهاتف'),
    'phone_country': ('Phone Country', 'بلد الهاتف'),
    'asset_type': ('Asset Type', 'Device Type', 'نوع الجهاز', 'نوع الفلتر'),
    'serial_number': ('Serial Number', 'Serial', 'الرقم التسلسلي'),
    'manufacturer': ('Manufacturer', 'الشركة المصنعة'),
    'model': ('Model', 'الموديل'),
    'country': ('Country', 'الدولة', 'البلد'),
    'state': ('Governorate / State', 'Governorate', 'State', 'Province', 'المحافظة'),
    'area': ('Area / Locality', 'Area', 'Locality', 'المنطقة'),
    'address': ('Address', 'العنوان'),
    'location_url': ('Location URL', 'Location', 'رابط الموقع'),
    'installation_date': ('Installation Date', 'Install Date', 'تاريخ التركيب'),
    'maintenance_cycle': ('Maintenance Cycle (Months)', 'Maintenance Cycle', 'دورة الصيانة'),
    'warranty_years': ('Warranty Years', 'Warranty Period (Years)', 'سنوات الضمان'),
    'status': ('Status', 'الحالة'),
    'notes': ('Notes', 'ملاحظات'),
}
MODES = {'customers', 'assets', 'combined'}


def heading(value):
    return re.sub(r'[\s_()/–-]+', ' ', value.strip().casefold()).strip()


ALIASES = {heading(alias): key for key, aliases in FIELDS.items() for alias in (*aliases, key)}


def auto_mapping(columns):
    seen, result = set(), []
    for column in columns:
        key = ALIASES.get(heading(column), '')
        result.append(key if key not in seen else '')
        if key: seen.add(key)
    return result


def file_error(message):
    raise ValueError(message)


def parse_file(encoded, format):
    if not isinstance(format, str) or format not in {'csv', 'xlsx'} or not isinstance(encoded, str) or len(encoded) > 4 * ((MAX_BYTES + 2) // 3):
        file_error('Choose an XLSX or UTF-8 CSV file of 1 MB or smaller.')
    try:
        content = base64.b64decode(encoded, validate=True)
        if not content or len(content) > MAX_BYTES: file_error('File must be nonempty and at most 1 MB.')
        if format == 'csv':
            text = content.decode('utf-8-sig')
            if '\x00' in text: file_error('CSV contains unsupported binary content.')
            rows = []
            for row in csv.reader(io.StringIO(text, newline=''), strict=True):
                rows.append(row)
                if len(rows) > MAX_ROWS + 1: file_error('Use at most 1000 data rows.')
        else:
            # Preflight compressed and expanded size, XML entities, active content and links.
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 100 or sum(e.file_size for e in entries) > 16 * MAX_BYTES:
                    file_error('Workbook expanded size exceeds the safe limit.')
                if len({e.filename for e in entries}) != len(entries): file_error('Workbook has duplicate archive entries.')
                for entry in entries:
                    name = entry.filename.lower()
                    if entry.flag_bits & 1 or '..' in name or name.startswith('/'):
                        file_error('Unsupported workbook archive.')
                    if any(part in name for part in ('vbaproject', 'externallinks', 'embeddings/', 'activex/', 'connections.xml')):
                        file_error('Macros, embedded objects and external workbook content are not supported.')
                    if name.endswith(('.xml', '.rels')):
                        root = fromstring(archive.read(entry))
                        for element in root.iter():
                            tag = element.tag.rsplit('}', 1)[-1]
                            if any('macroenabled' in value.casefold() for value in element.attrib.values()):
                                file_error('Macro-enabled workbooks are not supported. Save as a clean XLSX file.')
                            if tag in {'f', 'definedName'}: file_error('Formulas and defined names are not supported. Paste values into a clean workbook.')
                            if tag == 'Relationship' and element.attrib.get('TargetMode') == 'External':
                                file_error('External workbook links are not supported. Remove links before uploading.')
                            if tag == 'row' and int(element.attrib.get('r', '0')) > MAX_ROWS + 1:
                                file_error('Use at most 1000 data rows.')
                            if tag == 'c':
                                ref = element.attrib.get('r', '')
                                letters = re.match(r'[A-Z]+', ref)
                                if letters:
                                    col = 0
                                    for letter in letters[0]: col = col * 26 + ord(letter) - 64
                                    if col > MAX_COLUMNS: file_error('Use at most 40 columns.')
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) != 1: file_error('Use a workbook with one worksheet.')
                sheet = workbook.worksheets[0]
                sheet.reset_dimensions()
                rows = []
                for cells in sheet.iter_rows():
                    values = []
                    for cell in cells:
                        if cell.data_type in {'f', 'e'}: file_error('Formula and error cells are not supported.')
                        value = cell.value
                        if isinstance(value, datetime):
                            if value.time() != datetime.min.time(): file_error('Installation dates must not include a time.')
                            value = value.date()
                        if isinstance(value, date): value = value.isoformat()
                        if isinstance(value, bool): file_error('Boolean cells are not supported.')
                        if isinstance(value, float) and value.is_integer(): value = int(value)
                        values.append('' if value is None else str(value))
                    rows.append(values)
                    if len(rows) > MAX_ROWS + 1: file_error('Use at most 1000 data rows.')
            finally: workbook.close()
        while rows and not any(value.strip() for value in rows[-1]): rows.pop()
        if len(rows) < 2: file_error('Include a header and at least one data row.')
        columns = rows[0]
        if not columns or len(columns) > MAX_COLUMNS or any(not c.strip() or len(c) > 100 for c in columns):
            file_error('Use 1–40 nonempty column headings of at most 100 characters.')
        if len(set(map(heading, columns))) != len(columns): file_error('Column headings must be unique.')
        for index, row in enumerate(rows[1:], 2):
            # XLSX may omit empty trailing cells; CSV must keep its column count.
            if format == 'xlsx' and len(row) < len(columns): row.extend([''] * (len(columns) - len(row)))
            if len(row) != len(columns): file_error(f'Row {index} has a different number of columns than its headings.')
            if any(len(c) > MAX_CELL or '\x00' in c for c in row): file_error(f'Row {index} contains a cell longer than 1000 characters or unsupported content.')
            if any(c.lstrip().startswith('=') for c in row): file_error(f'Row {index} contains a formula. Paste plain values instead.')
        return {'columns': columns, 'rows': rows[1:], 'mapping': auto_mapping(columns), 'fields': {key: values[0] for key, values in FIELDS.items()}}
    except ValueError: raise
    except Exception:
        file_error('Malformed or unsupported spreadsheet. Use a clean XLSX or UTF-8 CSV file.')


def template(mode, format):
    if mode not in MODES or format not in {'csv', 'xlsx'}: file_error('Choose a supported import type and template format.')
    keys = list(FIELDS) if mode != 'customers' else ['customer_name', 'phone', 'phone_label', 'phone_country', 'country', 'status']
    labels = [FIELDS[key][0] for key in keys]
    if format == 'csv':
        output = io.StringIO(newline=''); csv.writer(output).writerow(labels)
        content = output.getvalue().encode('utf-8-sig')
    else:
        workbook = Workbook(); sheet = workbook.active; sheet.title = 'Import Data'
        sheet.append(labels); sheet.freeze_panes = 'A2'
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
        for index, key in enumerate(keys, 1):
            sheet.cell(1, index).font = Font(bold=True)
            sheet.column_dimensions[get_column_letter(index)].width = min(36, max(20, len(FIELDS[key][0]) + 3))
            # Empty cells formatted as text preserve leading phone zeroes and dates.
            for row in range(2, MAX_ROWS + 2): sheet.cell(row, index).number_format = '@'
            if key == 'status':
                validation = DataValidation(type='list', formula1='"Active,Inactive"'); validation.error = 'Choose Active or Inactive.'; validation.showErrorMessage = True
                sheet.add_data_validation(validation); validation.add(f'{get_column_letter(index)}2:{get_column_letter(index)}1001')
        output = io.BytesIO(); workbook.save(output); content = output.getvalue()
    return {'format': format, 'content': base64.b64encode(content).decode('ascii'), 'filename': f'axyrel-{mode}-template.{format}', 'max_bytes': MAX_BYTES, 'max_rows': MAX_ROWS}
